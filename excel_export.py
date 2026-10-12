import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import get_connection

def apply_header_style(cell):
    cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='15803D', end_color='15803D', fill_type='solid') # Forest green
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def apply_title_style(ws, title_text, col_span=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_span)
    cell = ws.cell(row=1, column=1)
    cell.value = title_text
    cell.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='166534', end_color='166534', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36

def apply_summary_style(cell, is_highlight=False):
    cell.font = Font(name='Calibri', size=11, bold=True, color='0F172A')
    fill_color = 'DCFCE7' if is_highlight else 'F1F5F9'
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')

def autofit_columns(ws, min_col=1, max_col=None):
    if max_col is None:
        max_col = ws.max_column
    for col in range(min_col, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def export_purchases_excel():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT p.bill_no, p.purchase_date, p.supplier_name, s.mobile, s.village,
               p.variety, p.quantity, p.unit, p.rate_per_unit, p.total_amount,
               p.paid_amount, p.remaining_amount, p.payment_status, p.vehicle_no,
               p.weighbridge_slip_no, p.notes
        FROM purchases p
        LEFT JOIN suppliers s ON p.supplier_id = s.id
        ORDER BY p.purchase_date DESC, p.id DESC
    ''')
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchases (Farmers)"
    ws.views.sheetView[0].showGridLines = True

    apply_title_style(ws, "CaneMart - Sugarcane Purchase Ledger (Farmer Procurement)", col_span=16)

    headers = [
        "Bill No", "Purchase Date", "Farmer Name", "Mobile", "Village",
        "Cane Variety", "Qty", "Unit", "Rate (₹/Unit)", "Total Amount (₹)",
        "Paid Amount (₹)", "Remaining (₹)", "Status", "Vehicle No", "Slip No", "Notes"
    ]

    ws.append([]) # blank row 2
    ws.append(headers) # row 3
    ws.row_dimensions[3].height = 26

    for col in range(1, len(headers) + 1):
        apply_header_style(ws.cell(row=3, column=col))

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    total_qty = 0
    total_val = 0
    total_paid = 0
    total_rem = 0

    current_row = 4
    for r in rows:
        data_row = [
            r['bill_no'], r['purchase_date'], r['supplier_name'], r['mobile'] or '-', r['village'] or '-',
            r['variety'], float(r['quantity']), r['unit'], float(r['rate_per_unit']), float(r['total_amount']),
            float(r['paid_amount']), float(r['remaining_amount']), r['payment_status'],
            r['vehicle_no'] or '-', r['weighbridge_slip_no'] or '-', r['notes'] or ''
        ]
        ws.append(data_row)
        ws.row_dimensions[current_row].height = 20

        total_qty += float(r['quantity'])
        total_val += float(r['total_amount'])
        total_paid += float(r['paid_amount'])
        total_rem += float(r['remaining_amount'])

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=c)
            cell.border = thin_border
            if c in [7, 9, 10, 11, 12]:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif c in [1, 2, 8, 13]:
                cell.alignment = Alignment(horizontal='center')

        current_row += 1

    # Add Summary Row
    summary_row = [
        "TOTAL", "", "", "", "", "", total_qty, "", "", total_val,
        total_paid, total_rem, "", "", "", ""
    ]
    ws.append(summary_row)
    ws.row_dimensions[current_row].height = 24
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=current_row, column=c)
        apply_summary_style(cell, is_highlight=(c in [10, 12]))
        cell.border = thin_border
        if c in [7, 10, 11, 12]:
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right')
        elif c == 1:
            cell.alignment = Alignment(horizontal='center')

    autofit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_sales_excel():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT s.invoice_no, s.sale_date, s.customer_name, c.mobile, c.location,
               c.customer_type, s.variety, s.purpose, s.quantity, s.unit,
               s.rate_per_unit, s.total_amount, s.paid_amount, s.remaining_amount,
               s.payment_status, s.vehicle_no, s.notes
        FROM sales s
        LEFT JOIN customers c ON s.customer_id = c.id
        ORDER BY s.sale_date DESC, s.id DESC
    ''')
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales (Animal Feed & Buyers)"
    ws.views.sheetView[0].showGridLines = True

    apply_title_style(ws, "CaneMart - Sugarcane Sales & Distribution Ledger (Animal Feed / Mills)", col_span=17)

    headers = [
        "Invoice No", "Sale Date", "Customer / Buyer", "Mobile", "Location",
        "Category", "Variety", "Purpose", "Qty", "Unit",
        "Rate (₹/Unit)", "Total Amount (₹)", "Paid (₹)", "Remaining (₹)",
        "Status", "Vehicle No", "Notes"
    ]

    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 26

    for col in range(1, len(headers) + 1):
        apply_header_style(ws.cell(row=3, column=col))

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    total_qty = 0
    total_val = 0
    total_paid = 0
    total_rem = 0

    current_row = 4
    for r in rows:
        data_row = [
            r['invoice_no'], r['sale_date'], r['customer_name'], r['mobile'] or '-', r['location'] or '-',
            r['customer_type'] or '-', r['variety'], r['purpose'] or 'Animal Feed',
            float(r['quantity']), r['unit'], float(r['rate_per_unit']), float(r['total_amount']),
            float(r['paid_amount']), float(r['remaining_amount']), r['payment_status'],
            r['vehicle_no'] or '-', r['notes'] or ''
        ]
        ws.append(data_row)
        ws.row_dimensions[current_row].height = 20

        total_qty += float(r['quantity'])
        total_val += float(r['total_amount'])
        total_paid += float(r['paid_amount'])
        total_rem += float(r['remaining_amount'])

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=c)
            cell.border = thin_border
            if c in [9, 11, 12, 13, 14]:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif c in [1, 2, 10, 15]:
                cell.alignment = Alignment(horizontal='center')

        current_row += 1

    summary_row = [
        "TOTAL", "", "", "", "", "", "", "", total_qty, "",
        "", total_val, total_paid, total_rem, "", "", ""
    ]
    ws.append(summary_row)
    ws.row_dimensions[current_row].height = 24
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=current_row, column=c)
        apply_summary_style(cell, is_highlight=(c in [12, 14]))
        cell.border = thin_border
        if c in [9, 12, 13, 14]:
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right')
        elif c == 1:
            cell.alignment = Alignment(horizontal='center')

    autofit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_full_canemart_report():
    conn = get_connection()
    cursor = conn.cursor()

    # 1. Dashboard summary
    cursor.execute('''
        SELECT 
            COALESCE(SUM(quantity), 0) AS total_purchase_qty,
            COALESCE(SUM(total_amount), 0) AS total_purchase_amt,
            COALESCE(SUM(paid_amount), 0) AS total_purchase_paid,
            COALESCE(SUM(remaining_amount), 0) AS total_purchase_rem
        FROM purchases
    ''')
    pur_summary = cursor.fetchone()

    cursor.execute('''
        SELECT 
            COALESCE(SUM(quantity), 0) AS total_sale_qty,
            COALESCE(SUM(total_amount), 0) AS total_sale_amt,
            COALESCE(SUM(paid_amount), 0) AS total_sale_paid,
            COALESCE(SUM(remaining_amount), 0) AS total_sale_rem
        FROM sales
    ''')
    sale_summary = cursor.fetchone()

    # Create Master Workbook
    wb = Workbook()

    # Sheet 1: Executive Summary
    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.views.sheetView[0].showGridLines = True
    apply_title_style(ws1, "CaneMart Business Overview & Financial Summary", col_span=4)

    summary_cards = [
        ("Key Business Metric", "Value", "Unit / Status", "Notes"),
        ("Total Sugarcane Purchased", pur_summary['total_purchase_qty'], "Tons", "Procured from farmers"),
        ("Total Sugarcane Sold / Dispatched", sale_summary['total_sale_qty'], "Tons", "Supplied to buyers & dairy"),
        ("Current Sugarcane Stock in Yard", pur_summary['total_purchase_qty'] - sale_summary['total_sale_qty'], "Tons", "Net available stock"),
        ("Total Purchase Expense (Cost)", pur_summary['total_purchase_amt'], "₹ INR", "Payable to supplier farmers"),
        ("Total Sales Revenue (Turnover)", sale_summary['total_sale_amt'], "₹ INR", "Billed to customers"),
        ("Gross Trading Margin / Profit", sale_summary['total_sale_amt'] - pur_summary['total_purchase_amt'], "₹ INR", "Turnover minus Cane Cost"),
        ("Pending Payable to Farmers", pur_summary['total_purchase_rem'], "₹ INR", "Liability to be paid"),
        ("Pending Receivable from Customers", sale_summary['total_sale_rem'], "₹ INR", "Assets to be collected")
    ]

    ws1.append([])
    for idx, row in enumerate(summary_cards):
        ws1.append(list(row))
        row_idx = idx + 3
        ws1.row_dimensions[row_idx].height = 24
        if idx == 0:
            for c in range(1, 5):
                apply_header_style(ws1.cell(row=row_idx, column=c))
        else:
            for c in range(1, 5):
                cell = ws1.cell(row=row_idx, column=c)
                if c == 2 and isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal='right')
                    apply_summary_style(cell, is_highlight=(idx in [4, 6, 7]))
                elif c == 1:
                    cell.font = Font(name='Calibri', size=11, bold=True)
    autofit_columns(ws1)

    # Sheet 2: Suppliers Ledger
    cursor.execute('''
        SELECT s.id, s.name, s.mobile, s.village, s.cane_variety, s.farm_area_acres,
               COALESCE(SUM(p.quantity), 0) AS total_qty,
               COALESCE(SUM(p.total_amount), 0) AS total_val,
               COALESCE(SUM(p.paid_amount), 0) AS paid_val,
               COALESCE(SUM(p.remaining_amount), 0) AS pending_val
        FROM suppliers s
        LEFT JOIN purchases p ON s.id = p.supplier_id
        GROUP BY s.id
        ORDER BY s.name ASC
    ''')
    suppliers = cursor.fetchall()

    ws2 = wb.create_sheet(title="Suppliers (Farmers)")
    ws2.views.sheetView[0].showGridLines = True
    apply_title_style(ws2, "CaneMart - Suppliers / Cane Farmers Directory & Balances", col_span=10)
    ws2.append([])
    headers2 = ["ID", "Farmer Name", "Mobile", "Village", "Cane Variety", "Acres", "Total Qty (Tons)", "Total Billed (₹)", "Paid (₹)", "Pending Balance (₹)"]
    ws2.append(headers2)
    ws2.row_dimensions[3].height = 26
    for c in range(1, 11):
        apply_header_style(ws2.cell(row=3, column=c))
    for r in suppliers:
        ws2.append([
            r['id'], r['name'], r['mobile'] or '-', r['village'] or '-', r['cane_variety'] or '-',
            float(r['farm_area_acres'] or 0), float(r['total_qty']), float(r['total_val']), float(r['paid_val']), float(r['pending_val'])
        ])
    autofit_columns(ws2)

    # Sheet 3: Customers Ledger
    cursor.execute('''
        SELECT c.id, c.name, c.mobile, c.location, c.customer_type,
               COALESCE(SUM(s.quantity), 0) AS total_qty,
               COALESCE(SUM(s.total_amount), 0) AS total_val,
               COALESCE(SUM(s.paid_amount), 0) AS paid_val,
               COALESCE(SUM(s.remaining_amount), 0) AS pending_val
        FROM customers c
        LEFT JOIN sales s ON c.id = s.customer_id
        GROUP BY c.id
        ORDER BY c.name ASC
    ''')
    customers = cursor.fetchall()

    ws3 = wb.create_sheet(title="Customers (Buyers)")
    ws3.views.sheetView[0].showGridLines = True
    apply_title_style(ws3, "CaneMart - Customers / Animal Feed & Mills Ledger", col_span=10)
    ws3.append([])
    headers3 = ["ID", "Customer Name", "Mobile", "Location", "Category", "Total Qty (Tons)", "Total Sales (₹)", "Paid Received (₹)", "Pending Balance (₹)", "Notes"]
    ws3.append(headers3)
    ws3.row_dimensions[3].height = 26
    for c in range(1, 11):
        apply_header_style(ws3.cell(row=3, column=c))
    for r in customers:
        ws3.append([
            r['id'], r['name'], r['mobile'] or '-', r['location'] or '-', r['customer_type'] or '-',
            float(r['total_qty']), float(r['total_val']), float(r['paid_val']), float(r['pending_val']), ""
        ])
    autofit_columns(ws3)

    # Sheet 4: Purchases Detail
    cursor.execute('SELECT * FROM purchases ORDER BY purchase_date DESC')
    purchases_all = cursor.fetchall()
    ws4 = wb.create_sheet(title="Purchase Records")
    ws4.views.sheetView[0].showGridLines = True
    apply_title_style(ws4, "All Sugarcane Purchase Transactions", col_span=11)
    ws4.append([])
    h4 = ["Bill No", "Date", "Farmer Name", "Variety", "Qty", "Unit", "Rate (₹)", "Total (₹)", "Paid (₹)", "Balance (₹)", "Status"]
    ws4.append(h4)
    for c in range(1, 12):
        apply_header_style(ws4.cell(row=3, column=c))
    for p in purchases_all:
        ws4.append([p['bill_no'], p['purchase_date'], p['supplier_name'], p['variety'], float(p['quantity']), p['unit'], float(p['rate_per_unit']), float(p['total_amount']), float(p['paid_amount']), float(p['remaining_amount']), p['payment_status']])
    autofit_columns(ws4)

    # Sheet 5: Sales Detail
    cursor.execute('SELECT * FROM sales ORDER BY sale_date DESC')
    sales_all = cursor.fetchall()
    ws5 = wb.create_sheet(title="Sales Records")
    ws5.views.sheetView[0].showGridLines = True
    apply_title_style(ws5, "All Sugarcane Sale Transactions", col_span=12)
    ws5.append([])
    h5 = ["Invoice No", "Date", "Customer Name", "Variety", "Purpose", "Qty", "Unit", "Rate (₹)", "Total (₹)", "Paid (₹)", "Balance (₹)", "Status"]
    ws5.append(h5)
    for c in range(1, 13):
        apply_header_style(ws5.cell(row=3, column=c))
    for s in sales_all:
        ws5.append([s['invoice_no'], s['sale_date'], s['customer_name'], s['variety'], s['purpose'], float(s['quantity']), s['unit'], float(s['rate_per_unit']), float(s['total_amount']), float(s['paid_amount']), float(s['remaining_amount']), s['payment_status']])
    autofit_columns(ws5)

    conn.close()

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
